"""Regression tests for offline analytics and report exports."""
from reporting import export_bundle


def test_report_bundle_contains_required_formats():
    rows = [{
        "event_id": "EVT-1",
        "timestamp": "2026-09-25T10:00:00+00:00",
        "severity": "CRITICAL",
        "source_ip": "10.0.0.1",
        "target_endpoint": "/login",
        "attack_type": "SQL injection",
        "description": "test",
        "raw_sha256": "a" * 64,
    }]
    analysis = {
        "risk_score": 80,
        "alerts": [object()],
        "analyst_alerts": [{
            "alert_id": "FND-000001", "rule_id": "SQL_INJECTION", "rule_version": "1.0",
            "title": "SQL injection indicator", "severity": "CRITICAL", "confidence": "HIGH",
            "category": "WEB_ATTACK", "mitre_technique": "T1190", "count": 1,
            "sources": ["10.0.0.1"], "destinations": ["/login"],
            "first_seen": "2026-09-25T10:00:00+00:00", "last_seen": "2026-09-25T10:00:00+00:00",
            "guidance": "review",
        }],
        "incidents": [],
        "parse_coverage": 100.0,
    }
    bundle = export_bundle(rows, analysis, "test")
    assert {"pdf", "docx", "xlsx", "pptx", "html", "csv", "json"} <= set(bundle)
    for payload in bundle.values():
        assert isinstance(payload, bytes)
        assert payload

def test_xlsx_preserves_all_evidence_fields_and_records_and_is_print_ready():
    rows = [
        {"event_id": "EVT-1", "timestamp": "2026-09-25T10:00:00+00:00", "custom_field": "KEEP-ME", "nested_text": "{\"x\":1}"},
        {"event_id": "EVT-2", "timestamp": "2026-09-25T10:01:00+00:00", "another_field": "SECOND"},
    ]
    payload = export_bundle(rows, {"risk_score": 10, "alerts": [], "analyst_alerts": [], "incidents": [], "parse_coverage": 100}, "test")["xlsx"]
    from io import BytesIO
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(payload), data_only=False)
    assert "Evidence" in wb.sheetnames
    ws = wb["Evidence"]
    headers = [c.value for c in ws[1]]
    assert headers == ["event_id", "timestamp", "custom_field", "nested_text", "another_field"]
    assert ws.max_row == 3
    assert ws["C2"].value == "KEEP-ME"
    assert ws["E3"].value == "SECOND"
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == ws.dimensions
    assert ws.page_setup.paperSize == ws.PAPERSIZE_A4
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.fitToWidth == 1
    assert ws.print_area == ws.dimensions
    assert ws.oddFooter.center.text == "VANGUARD SOC • Page &P of &N"
