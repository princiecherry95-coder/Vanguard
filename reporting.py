"""Offline multi-format SOC report generation. No network dependencies."""
from __future__ import annotations
from io import BytesIO
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
import json
import pandas as pd

def _safe_name(name: str) -> str:
    return "".join(c if c.isalnum() or c in "-_" else "_" for c in name)[:80] or "vanguard_report"

def report_context(rows: list[dict[str,Any]], analysis: dict[str,Any] | None, source: str) -> dict[str,Any]:
    from analytics import summary
    s=summary(rows)
    return {"system":"Vanguard-SIEM","classification":"LOCAL DEFENSIVE ANALYTICS","source":source,
            "generated_at":datetime.now(timezone.utc).isoformat(),"analytics":s,
            "risk_score":(analysis or {}).get("risk_score",0),
            "raw_alerts":len((analysis or {}).get("alerts",[])),
            "finding_groups":len((analysis or {}).get("analyst_alerts",[])),
            "incidents":len((analysis or {}).get("incidents",[]))}

def make_pdf(rows, analysis, source, title="Vanguard-SIEM SOC Analytics Report") -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    ctx=report_context(rows,analysis,source); s=ctx["analytics"]; styles=getSampleStyleSheet()
    out=BytesIO(); doc=SimpleDocTemplate(out,pagesize=A4,rightMargin=14*mm,leftMargin=14*mm,topMargin=14*mm,bottomMargin=14*mm)
    story=[Paragraph(title,styles["Title"]),Paragraph(f"Source: {source}<br/>Generated: {ctx['generated_at']}",styles["Normal"]),Spacer(1,8)]
    story.append(Table([["Metric","Value"],["Records",s["records"]],["Unique sources",s["unique_sources"]],["Unique targets",s["unique_targets"]],["Risk score",ctx["risk_score"]],["Raw alerts",ctx["raw_alerts"]],["Finding groups",ctx["finding_groups"]],["Incidents",ctx["incidents"]]],repeatRows=1))
    story += [Spacer(1,10),Paragraph("Severity distribution",styles["Heading2"])]
    story.append(Table([["Severity","Count"]]+[[k,v] for k,v in s["severity_counts"].items()],repeatRows=1))
    story += [Spacer(1,10),Paragraph("Top attack types",styles["Heading2"])]
    story.append(Table([["Attack type","Count"]]+[[str(k),v] for k,v in s["top_attack_types"].items()],repeatRows=1))
    story += [Spacer(1,10),Paragraph("Top sources",styles["Heading2"])]
    story.append(Table([["Source IP","Events"]]+[[str(k),v] for k,v in s["top_sources"].items()],repeatRows=1))
    story += [Spacer(1,10),Paragraph("Evidence records",styles["Heading2"])]
    data=[["Event ID","Time","Severity","Source","Attack type"]]
    for r in rows[:500]:
        data.append([str(r.get("event_id","")),str(r.get("timestamp","")),str(r.get("severity","")),str(r.get("source_ip","")),str(r.get("attack_type",""))[:70]])
    t=Table(data,repeatRows=1)
    t.setStyle(TableStyle([("GRID",(0,0),(-1,-1),.25,colors.grey),("FONTSIZE",(0,0),(-1,-1),6),("VALIGN",(0,0),(-1,-1),"TOP")]))
    story += [t]; doc.build(story); return out.getvalue()

def make_docx(rows, analysis, source, title="Vanguard-SIEM SOC Analytics Report") -> bytes:
    from docx import Document
    d=Document(); ctx=report_context(rows,analysis,source); s=ctx["analytics"]
    d.add_heading(title,0); d.add_paragraph(f"Source: {source}\nGenerated: {ctx['generated_at']}")
    d.add_heading("Executive metrics",1)
    table=d.add_table(rows=1, cols=2); table.rows[0].cells[0].text="Metric"; table.rows[0].cells[1].text="Value"
    for k,v in [("Records",s["records"]),("Unique sources",s["unique_sources"]),("Unique targets",s["unique_targets"]),("Risk score",ctx["risk_score"]),("Raw alerts",ctx["raw_alerts"]),("Finding groups",ctx["finding_groups"]),("Incidents",ctx["incidents"])]:
        c=table.add_row().cells; c[0].text=k; c[1].text=str(v)
    d.add_heading("Severity distribution",1)
    for k,v in s["severity_counts"].items(): d.add_paragraph(f"{k}: {v}")
    d.add_heading("Top attack types",1)
    for k,v in s["top_attack_types"].items(): d.add_paragraph(f"{k}: {v}")
    d.add_heading("Evidence sample",1)
    t=d.add_table(rows=1,cols=5)
    for c,h in zip(t.rows[0].cells,["Event ID","Time","Severity","Source","Attack"]): c.text=h
    for r in rows[:500]:
        c=t.add_row().cells
        vals=[r.get("event_id",""),r.get("timestamp",""),r.get("severity",""),r.get("source_ip",""),r.get("attack_type","")]
        for cell,val in zip(c,vals): cell.text=str(val)[:100]
    out=BytesIO(); d.save(out); return out.getvalue()

def make_xlsx(rows, analysis, source) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.utils import get_column_letter
    ctx=report_context(rows,analysis,source); s=ctx["analytics"]; wb=Workbook(); ws=wb.active; ws.title="Summary"
    metrics=[("Source",source),("Generated",ctx["generated_at"]),("Records",s["records"]),("Unique sources",s["unique_sources"]),("Unique targets",s["unique_targets"]),("Risk score",ctx["risk_score"]),("Raw alerts",ctx["raw_alerts"]),("Finding groups",ctx["finding_groups"]),("Incidents",ctx["incidents"])]
    for i,(k,v) in enumerate(metrics,1): ws.cell(i,1,k); ws.cell(i,2,v)
    sev=wb.create_sheet("Severity"); sev.append(["Severity","Count"]); [sev.append([k,v]) for k,v in s["severity_counts"].items()]
    chart=BarChart(); chart.title="Severity"; chart.y_axis.title="Events"; chart.x_axis.title="Severity"; chart.add_data(Reference(sev,min_col=2,min_row=1,max_row=1+len(s["severity_counts"])),titles_from_data=True); chart.set_categories(Reference(sev,min_col=1,min_row=2,max_row=1+len(s["severity_counts"]))); sev.add_chart(chart,"D2")
    ev=wb.create_sheet("Evidence"); cols=sorted({k for r in rows[:5000] for k in r}) if rows else ["event_id"]; ev.append(cols)
    for r in rows[:5000]: ev.append([r.get(c,"") for c in cols])
    for sheet in wb.worksheets:
        for col in range(1,sheet.max_column+1): sheet.column_dimensions[get_column_letter(col)].width=min(45,max(12,max((len(str(sheet.cell(row=row,column=col).value or "")) for row in range(1,min(sheet.max_row,100)+1)),default=12)+2))
    out=BytesIO(); wb.save(out); return out.getvalue()

def make_pptx(rows, analysis, source, title="Vanguard-SIEM SOC Analytics Report") -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    ctx=report_context(rows,analysis,source); s=ctx["analytics"]; prs=Presentation()
    slide=prs.slides.add_slide(prs.slide_layouts[0]); slide.shapes.title.text=title; slide.placeholders[1].text=f"{source}\n{ctx['generated_at']}"
    slide=prs.slides.add_slide(prs.slide_layouts[5]); slide.shapes.title.text="Executive analytics"
    box=slide.shapes.add_textbox(Inches(1),Inches(1.5),Inches(11),Inches(4)).text_frame
    box.text="Records: %s\nRisk score: %s\nRaw alerts: %s\nFinding groups: %s\nIncidents: %s\nUnique sources: %s\nUnique targets: %s" % (s["records"],ctx["risk_score"],ctx["raw_alerts"],ctx["finding_groups"],ctx["incidents"],s["unique_sources"],s["unique_targets"])
    slide=prs.slides.add_slide(prs.slide_layouts[5]); slide.shapes.title.text="Severity distribution"
    tb=slide.shapes.add_textbox(Inches(1),Inches(1.4),Inches(11),Inches(4)).text_frame; tb.text="\n".join(f"{k}: {v}" for k,v in s["severity_counts"].items())
    slide=prs.slides.add_slide(prs.slide_layouts[5]); slide.shapes.title.text="Top attack types"
    tb=slide.shapes.add_textbox(Inches(1),Inches(1.4),Inches(11),Inches(5)).text_frame; tb.text="\n".join(f"{k}: {v}" for k,v in s["top_attack_types"].items()) or "No detected attack types."
    out=BytesIO(); prs.save(out); return out.getvalue()

def export_bundle(rows, analysis, source, basename="vanguard_soc_report") -> dict[str,bytes]:
    return {"pdf":make_pdf(rows,analysis,source),"docx":make_docx(rows,analysis,source),"xlsx":make_xlsx(rows,analysis,source),"pptx":make_pptx(rows,analysis,source),
            "json":json.dumps(report_context(rows,analysis,source),indent=2).encode()}
